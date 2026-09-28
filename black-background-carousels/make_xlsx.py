#!/usr/bin/env python3
"""Build the week-4 batch workbook from batch_week4.py."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from batch_week4 import build_rows, CLOSER
from captions import all_captions

PINK = "FF6B9D"
DARK = "1A1416"
BAND = "FFF3F7"

rows = build_rows(start_id=43)
CAPS = all_captions(43)

wb = Workbook()

# ---------------------------------------------------------------- Batch sheet
ws = wb.active
ws.title = "Batch"

headers = ["ID", "Feature", "Aimed at", "Slide 1 — claim", "Slide 1 — undercut",
           "Slide 2 — line A", "Slide 2 — line B", "Slide 3 — line A",
           "Slide 3 — line B", "Comment keyword", "Slide 4 — payoff",
           "Closer", "Caption", "Keep?", "Notes"]
ws.append(headers)

for r in rows:
    ws.append([r["id"], r["feature"], r["aimed_at"], r["s1a"], r["s1b"],
               r["s2a"], r["s2b"], r["s3a"], r["s3b"], r["keyword"],
               r["s4"], r["closer"], CAPS[r["id"]], "", ""])

head_fill = PatternFill("solid", fgColor=DARK)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.fill = head_fill
    cell.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 34
ws.freeze_panes = "D2"

thin = Side(style="thin", color="E8DDE3")
band = PatternFill("solid", fgColor=BAND)
for i in range(2, len(rows) + 2):
    feature = ws.cell(row=i, column=2).value
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=i, column=c)
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
        if feature == "Brownies":
            cell.fill = band
    ws.cell(row=i, column=1).font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=i, column=10).font = Font(name="Arial", size=10, bold=True,
                                          color=PINK)

widths = [8, 11, 10, 30, 30, 26, 26, 26, 26, 14, 26, 30, 52, 8, 26]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

dv = DataValidation(type="list", formula1='"yes,no,maybe"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"N2:N{len(rows) + 1}")

# ------------------------------------------------------------- Summary sheet
s = wb.create_sheet("Summary")
s["A1"] = "CoupleIn — week 4 batch"
s["A1"].font = Font(name="Arial", size=14, bold=True, color=DARK)
s["A2"] = "90 carousels, P043–P132. Four slides each."
s["A2"].font = Font(name="Arial", size=10, color="7A6E70")

s["A4"] = "Feature"
s["B4"] = "Posts"
s["C4"] = "Aimed at him"
s["D4"] = "Aimed at her"
for c in "ABCD":
    s[f"{c}4"].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    s[f"{c}4"].fill = head_fill

last = len(rows) + 1
for i, feature in enumerate(["Calendar", "Brownies", "Resolve"], start=5):
    s[f"A{i}"] = feature
    s[f"B{i}"] = f'=COUNTIF(Batch!$B$2:$B${last},$A{i})'
    s[f"C{i}"] = f'=COUNTIFS(Batch!$B$2:$B${last},$A{i},Batch!$C$2:$C${last},"he")'
    s[f"D{i}"] = f'=COUNTIFS(Batch!$B$2:$B${last},$A{i},Batch!$C$2:$C${last},"she")'
s["A8"] = "Total"
s["B8"] = "=SUM(B5:B7)"
s["C8"] = "=SUM(C5:C7)"
s["D8"] = "=SUM(D5:D7)"
for c in "ABCD":
    s[f"{c}8"].font = Font(name="Arial", size=10, bold=True)

s["A10"] = "Kept for production"
s["B10"] = f'=COUNTIF(Batch!$N$2:$N${last},"yes")'
s["A11"] = "Cut"
s["B11"] = f'=COUNTIF(Batch!$N$2:$N${last},"no")'
s["A12"] = "Not yet reviewed"
s["B12"] = f'=COUNTBLANK(Batch!$N$2:$N${last})'

notes = [
    "",
    "HOW TO USE THIS",
    "Mark column N on the Batch tab yes / no / maybe. The counts above update.",
    "Send me the sheet back and I'll render every 'yes' row at both sizes.",
    "",
    "THE FORMAT",
    "Slide 1  claim, then the everyday truth that undercuts it.",
    "Slide 2  the sharpening — twists the knife, stays concrete.",
    "Slide 3  the release: he's not refusing, he doesn't know how. Plus the comment prompt.",
    "Slide 4  app screenshot, one payoff line, then the closer.",
    "",
    "RULES THAT KEEP IT VIRAL-SHAPED",
    "Four to eight words a line. If a line needs a comma, it's too long.",
    "Slide 3 must release the sting, or the right move is to leave him, not download an app.",
    "Never use 'them' for one person — him and her throughout.",
    "",
    "FEATURE MAPPING",
    "Calendar   chores, forgetting, plans, the mental load.",
    "Brownies   appreciation, effort, what actually counts.",
    "Resolve    arguments, avoidance, shutting down.",
    "",
    "CLOSER ON EVERY POST",
    CLOSER,
]
r = 14
for line in notes:
    s[f"A{r}"] = line
    if line.isupper() and line:
        s[f"A{r}"].font = Font(name="Arial", size=10, bold=True, color=PINK)
    else:
        s[f"A{r}"].font = Font(name="Arial", size=10)
    r += 1

for col, w in zip("ABCD", [78, 12, 14, 14]):
    s.column_dimensions[col].width = w

for row in s.iter_rows(min_row=5, max_row=12, min_col=1, max_col=4):
    for cell in row:
        if cell.font.name != "Arial" or not cell.font.bold:
            cell.font = Font(name="Arial", size=10)

import os
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                     "CoupleIn_week4_90_hooks.xlsx"))
print("saved", len(rows), "rows")
